"""The enrichment pass, end to end after the database, with a stubbed provider.

Covers the configured and unconfigured states from the first commit: every
earlier "not configured collapsed into false" defect came from a state with no
test. The provider here is invented — the pass reads only what the `enrich`
package promises, so any provider behaves the same way to it.

No real Config(): it reads the environment, and a failing assertion can put a
live credential in a traceback. Stubs throughout.
"""

from __future__ import annotations

import json
import re
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from quorum import enrich
from quorum.crm.contact import Contact
from quorum.enrich import Company, Person
from quorum.weekly import coverage as coverage_mod
from quorum.weekly import enrichment
from quorum.weekly import people as people_mod
from quorum.weekly import stakeholders as stakeholders_mod
from quorum.weekly import view as view_mod
from quorum.weekly import workbook as workbook_mod

PROFILE = {
    "employee_count_min": 50,
    "employee_count_max": 5000,
    "hq_geographies": ["North America"],
    "focus_seniority": ["vp", "c-level"],
}

# What the CRM holds about each company. Globex is recorded at 10 employees —
# the misrecorded-headcount case that removes a company from the map unseen.
FIRMOGRAPHICS = {
    "acme.example": {"name": "Acme", "employees": 500, "country": "United States"},
    "globex.example": {"name": "Globex", "employees": 10, "country": "United States"},
    "initech.example": {"name": "Initech", "employees": 300, "country": "Canada"},
}

BENCH = {
    "acme.example": [
        Contact(name="Dana Reyes", title="VP Sales", email="dana@acme.example",
                linkedin="https://www.linkedin.com/in/dana-reyes"),
        Contact(name="Lee Park", title="VP Marketing", email="lee@acme.example"),
    ],
    # The CRM's LinkedIn URL for Kim points at someone else — the wrong-URL case.
    "globex.example": [Contact(name="Kim Lo", title="CRO", email="kim@globex.example",
                               linkedin="https://www.linkedin.com/in/kimlo")],
    # Ray's email finds nothing; his CRM LinkedIn URL finds him, somewhere else.
    "initech.example": [Contact(name="Ray Oh", title="VP Ops", email="ray@initech.example",
                                linkedin="https://linkedin.com/in/ray-oh/")],
}

IN_CRM = {"dana@acme.example", "lee@acme.example"}


class _SF:
    """Salesforce-only, the configuration the first deployment runs."""

    configured = True
    linkedin_available = True

    def contact_by_email(self, email):
        for bench in BENCH.values():
            for c in bench:
                if c.email == email and email in IN_CRM:
                    return c
        return None

    def domain_stats(self, domain, terms):
        return {"sf_total": 2, "sf_senior": 1, "account_id": domain}

    def account_firmographics(self, account_id):
        f = FIRMOGRAPHICS.get(account_id, {})
        return {"name": f.get("name", ""), "employees": f.get("employees", ""),
                "hq": f.get("country", ""), "country": f.get("country", ""),
                "account_type": "", "city": "", "state": ""}

    def senior_bench(self, domain, terms):
        return list(BENCH.get(domain, []))


class _HS:
    configured = False

    def contact_by_email(self, email):
        return None

    def count_domain(self, domain):
        return 0


class _Provider:
    """An invented provider with the interface the `enrich` package promises."""

    display_name = "Example"

    def __init__(self):
        self.person_calls: list[str] = []
        self.linkedin_calls: list[str] = []
        self.company_calls: list[str] = []
        self.people = {
            # Still at Acme, with a different title there — but an advisory
            # seat elsewhere is listed first. Still here, not moved.
            "dana@acme.example": Person(
                name="Dana Reyes", title="Advisor", employer_name="Board Co",
                employer_domain="board.example",
                linkedin="https://www.linkedin.com/in/dana-reyes",
                current_jobs=(("board.example", "Board Co", "Advisor"),
                              ("acme.example", "Acme", "SVP Sales")),
            ),
            # Moved on.
            "lee@acme.example": Person(name="Lee Park", title="CMO", employer_name="Globex",
                                       employer_domain="globex.example", updated="2026-06-01"),
            # Not in the CRM; the provider knows them.
            "ari@acme.example": Person(name="Ari Stone", title="Director of Ops",
                                       employer_name="Acme", employer_domain="acme.example",
                                       linkedin="https://www.linkedin.com/in/ari-stone"),
        }
        self.companies = {
            # Agrees with the CRM's verdict, though not its number.
            "acme.example": Company(name="Acme", domain="acme.example", employees=520,
                                    country="United States"),
            # Disagrees: the CRM says 10 employees, so the CRM says no.
            "globex.example": Company(name="Globex", domain="globex.example",
                                      employees=150, country="United States"),
        }

    def person_by_email(self, email):
        self.person_calls.append(email)
        return self.people.get(email)

    # The provider accepts on the handle alone; the name check is the pass's.
    by_linkedin = {
        "ray-oh": Person(name="Ray Oh", title="COO", employer_name="Hooli",
                         employer_domain="hooli.example",
                         linkedin="https://www.linkedin.com/in/ray-oh"),
        "kimlo": Person(name="Kimberly Stone", title="CRO", employer_name="Other Co",
                        employer_domain="other.example",
                        linkedin="https://www.linkedin.com/in/kimlo"),
    }

    def person_by_linkedin(self, url):
        from quorum.enrich import linkedin_handle

        self.linkedin_calls.append(url)
        return self.by_linkedin.get(linkedin_handle(url))

    def company_by_domain(self, domain):
        self.company_calls.append(domain)
        return self.companies.get(domain)


def _cfg():
    return SimpleNamespace(
        salesforce=SimpleNamespace(configured=True),
        hubspot=SimpleNamespace(configured=False),
        customer_account_types=(),
        recent_days=90,
        group_call_min=8,
        shortlist_size=3,
    )


def _attendees():
    rows = [
        {"attendee_name": "Dana Reyes", "email": "dana@acme.example", "domain": "acme.example"},
        {"attendee_name": "Lee Park", "email": "lee@acme.example", "domain": "acme.example"},
        {"attendee_name": "", "email": "ari@acme.example", "domain": "acme.example"},
        {"attendee_name": "Support", "email": "support@acme.example", "domain": "acme.example"},
        {"attendee_name": "Kim Lo", "email": "kim@globex.example", "domain": "globex.example"},
        {"attendee_name": "Pat Vo", "email": "pat@initech.example", "domain": "initech.example"},
    ]
    return people_mod.dedupe_people([{**r, "meeting_title": "Intro"} for r in rows])


def _run(tmp_path, provider):
    """Steps 3 to 6 in run.py's order, with the enrichment pass where run.py
    puts it: companies before the map is chosen, people after."""
    cfg, sf, hs = _cfg(), _SF(), _HS()
    people = _attendees()
    reconciled = [people_mod.reconcile(p, sf, hs) for p in people]
    coverage = coverage_mod.build_coverage(
        cfg, people_mod.group_companies(people), PROFILE, sf, hs
    )
    pass_ = enrichment.start(provider)
    queue = []
    if pass_:
        enrichment.companies(pass_, coverage, PROFILE)
    rows, _ = stakeholders_mod.build(
        cfg, coverage, coverage_mod.seniority_terms(PROFILE), {}, sf
    )
    if pass_:
        enrichment.stakeholders(pass_, rows)
        enrichment.not_in_crm(pass_, reconciled)
        queue = enrichment.review_queue(pass_, coverage, rows)

    xlsx = str(tmp_path / "weekly_stakeholder_map_2026-08-17.xlsx")
    workbook_mod.build_workbook(
        cfg, reconciled, coverage, [], rows, xlsx, profile=PROFILE,
        geo_label="North America",
        enrichment=provider.display_name if provider else None, queue=queue,
    )
    html = open(view_mod.render(xlsx, "example.com")).read()
    return load_workbook(xlsx), html, coverage


def _table(ws):
    """Header row and data rows as dicts, captions and blank rows dropped."""
    rows = list(ws.iter_rows(values_only=True))
    headers = [h for h in rows[0] if h]
    out = []
    for r in rows[1:]:
        # A caption sits alone in column A; a data row has something after it.
        # A blank Name is a real row — the case the provider's name fills.
        if any(v not in (None, "") for v in r[1:]):
            out.append(dict(zip(headers, r)))
    return headers, out


# --- off ---------------------------------------------------------------------- #


def test_with_no_provider_nothing_changes(tmp_path):
    wb, html, _ = _run(tmp_path, None)

    assert wb.sheetnames == [
        "1 - Met this week", "2 - Company coverage", "3 - Stakeholder list",
    ]
    for ws in wb.worksheets:
        headers, _ = _table(ws)
        assert not any("Example" in h or h in ("Profile check", "Still at company?")
                       for h in headers), ws.title
    # The CRM's "no" stands alone, so the misrecorded company is off the map —
    # the invisible error the pass exists to surface.
    _, rows = _table(wb["3 - Stakeholder list"])
    assert not any("Globex" in str(r["Company"]) for r in rows)
    assert "Review queue" not in html


# --- on ----------------------------------------------------------------------- #


def test_company_verdicts_are_compared_not_numbers(tmp_path):
    wb, _, _ = _run(tmp_path, _Provider())
    headers, rows = _table(wb["2 - Company coverage"])
    by = {r["Company"]: r for r in rows}

    assert headers[-3:] == ["Employees (Example)", "HQ (Example)", "Profile check"]
    # 500 against 520 is inside the band either way: not a finding.
    assert by["acme.example"]["Profile check"] == "agrees"
    assert by["globex.example"]["Profile check"] == "disputed — Example says yes"
    # Looked, and the provider has nothing. Stated, never blank.
    assert by["initech.example"]["Profile check"] == "not found in Example"
    assert by["initech.example"]["Employees (Example)"] is None
    # The CRM's own value is never replaced.
    assert by["globex.example"]["Employees"] == 10


def test_a_disputed_rejection_reaches_the_map_marked(tmp_path):
    wb, _, _ = _run(tmp_path, _Provider())
    _, rows = _table(wb["3 - Stakeholder list"])

    globex = [r for r in rows if str(r["Company"]).startswith("Globex")]
    assert globex and all(r["Company"] == "Globex (profile disputed)" for r in globex)
    # Agreeing companies are not marked.
    assert any(r["Company"] == "Acme" for r in rows)


def test_still_at_company_and_differences(tmp_path):
    wb, _, _ = _run(tmp_path, _Provider())
    headers, rows = _table(wb["3 - Stakeholder list"])
    by = {r["Name"]: r for r in rows}

    assert headers[-3:] == ["Still at company?", "Title (Example)", "LinkedIn (Example)"]
    # Acme is among Dana's current positions, though not listed first.
    assert by["Dana Reyes"]["Still at company?"] == "yes"
    assert by["Lee Park"]["Still at company?"] == "no — now at Globex, as CMO"
    # Kim's CRM LinkedIn URL points at someone else: not used, still not found.
    assert by["Kim Lo"]["Still at company?"] == "not found in Example"
    # Ray's email found nothing; his LinkedIn URL did, and the name agrees.
    assert by["Ray Oh"]["Still at company?"] == "no — now at Hooli, as COO (matched on LinkedIn)"
    # Shown only where it differs — and it is the title at *this* company, not
    # the first one listed. The CRM's title stays in Title.
    assert by["Dana Reyes"]["Title (Example)"] == "SVP Sales"
    assert by["Dana Reyes"]["Title"] == "VP Sales"
    assert by["Dana Reyes"]["LinkedIn (Example)"] in (None, "")
    # A detected move flags the row; it is not removed.
    assert "Lee Park" in by
    # Their provider title is for the new job, so it is not set beside this one.
    assert by["Lee Park"]["Title (Example)"] in (None, "")


def test_people_not_in_the_crm_get_a_name_title_and_linkedin(tmp_path):
    """On Met this week, for the people the CRM does not hold — including the
    LinkedIn URL, which is what someone needs to connect with them."""
    wb, html, _ = _run(tmp_path, _Provider())
    headers, rows = _table(wb["1 - Met this week"])
    by = {r["Email"]: r for r in rows}

    # The provider's name is in the header; the caption says what that means.
    assert "The Example columns are enrichment from Example, not your CRM" in html

    assert ["Name (Example)", "Title (Example)", "LinkedIn (Example)"] == [
        h for h in headers if "(Example)" in h
    ]
    assert by["ari@acme.example"]["Name (Example)"] == "Ari Stone"
    assert by["ari@acme.example"]["Title (Example)"] == "Director of Ops"
    assert by["ari@acme.example"]["LinkedIn (Example)"] == "https://www.linkedin.com/in/ari-stone"
    assert by["pat@initech.example"]["Name (Example)"] == "not found in Example"
    assert by["support@acme.example"]["Name (Example)"] == "not looked up — shared inbox"
    # People the CRM does hold are not given provider columns here: the
    # stakeholder list is where the provider is set beside a CRM record.
    assert by["dana@acme.example"]["Name (Example)"] in (None, "")


def test_the_summary_counts_what_the_provider_found_and_the_queue(tmp_path):
    """Ari is found; Kim and Pat are not; the shared inbox is never looked up,
    so it is in neither number."""
    from quorum.weekly import summary as summary_mod

    cfg, sf, hs = _cfg(), _SF(), _HS()
    people = _attendees()
    reconciled = [people_mod.reconcile(p, sf, hs) for p in people]
    coverage = coverage_mod.build_coverage(
        cfg, people_mod.group_companies(people), PROFILE, sf, hs
    )
    pass_ = enrichment.start(_Provider())
    enrichment.companies(pass_, coverage, PROFILE)
    rows, raw = stakeholders_mod.build(
        cfg, coverage, coverage_mod.seniority_terms(PROFILE), {}, sf
    )
    enrichment.stakeholders(pass_, rows)
    enrichment.not_in_crm(pass_, reconciled)
    queue = enrichment.review_queue(pass_, coverage, rows)

    stats = summary_mod.build(
        cfg, reconciled, coverage, rows, raw, PROFILE, enrichment="Example",
        queue=queue, queue_kinds=enrichment.queue_kinds(_Provider()),
    )
    by = {s["key"]: (s["count"], s["out_of"]) for s in stats}

    assert by["people_not_in_crm"] == (4, 6)
    assert by["people_not_in_crm_found"] == (1, 3)
    kinds = [s["key"][len("queue:"):] for s in stats if s["key"].startswith("queue:")]
    # The key is stable; only the label names the provider.
    assert kinds[: len(enrichment.QUEUE_ORDER)] == list(enrichment.QUEUE_ORDER)
    assert kinds[:2] == ["Profile fit disputed", "Company not found"]
    assert next(s["what"] for s in stats if s["key"] == "queue:Company not found") == (
        "Company not found in Example"
    )
    assert sum(s["count"] for s in stats if s["key"].startswith("queue:")) == len(queue)

    xlsx = str(tmp_path / "weekly_stakeholder_map_2026-08-17.xlsx")
    workbook_mod.build_workbook(
        cfg, reconciled, coverage, [], rows, xlsx, profile=PROFILE,
        geo_label="North America", enrichment="Example", queue=queue, summary=stats,
    )
    html = open(view_mod.render(xlsx, "example.com")).read()
    # The sheet keeps the full label; the page's line drops the words it just said.
    assert any("Not in your CRM, and found by Example" in str(c.value)
               for row in load_workbook(xlsx)["Summary"].iter_rows() for c in row)
    assert "Not in your CRM 4 of 6 · and found by Example" in html
    assert "Enrichment: Example — its findings are on the Review queue tab." in html


def test_with_no_provider_the_summary_says_so_in_words(tmp_path):
    """No column and no count is the rule for a source that was not asked. The
    Summary tab still says it in a sentence, because otherwise a run that was
    meant to be enriched and was not looks identical to one that never was."""
    from quorum.weekly import summary as summary_mod

    cfg, sf, hs = _cfg(), _SF(), _HS()
    people = _attendees()
    reconciled = [people_mod.reconcile(p, sf, hs) for p in people]
    coverage = coverage_mod.build_coverage(
        cfg, people_mod.group_companies(people), PROFILE, sf, hs
    )
    rows, raw = stakeholders_mod.build(
        cfg, coverage, coverage_mod.seniority_terms(PROFILE), {}, sf
    )
    stats = summary_mod.build(cfg, reconciled, coverage, rows, raw, PROFILE)

    assert not any(s["area"] == "Review queue" for s in stats)
    assert "people_not_in_crm_found" not in {s["key"] for s in stats}

    xlsx = str(tmp_path / "weekly_stakeholder_map_2026-08-17.xlsx")
    workbook_mod.build_workbook(
        cfg, reconciled, coverage, [], rows, xlsx, profile=PROFILE,
        geo_label="North America", summary=stats,
    )
    html = open(view_mod.render(xlsx, "example.com")).read()
    assert (
        "Enrichment: not configured — no provider lookups and no review queue this run."
        in html
    )
    assert "<h2>Review queue" not in html


def test_each_person_and_company_is_looked_up_once_and_inboxes_never(tmp_path):
    provider = _Provider()
    _run(tmp_path, provider)

    assert "support@acme.example" not in provider.person_calls
    assert len(provider.person_calls) == len(set(provider.person_calls))
    # LinkedIn is tried only where the email found nothing and the CRM holds a
    # URL: Kim and Ray here — never Dana or Lee, whose emails matched.
    assert sorted(provider.linkedin_calls) == [
        "https://linkedin.com/in/ray-oh/", "https://www.linkedin.com/in/kimlo",
    ]
    assert len(provider.company_calls) == len(set(provider.company_calls))


def test_the_review_queue_holds_what_a_person_should_settle(tmp_path):
    wb, html, _ = _run(tmp_path, _Provider())
    headers, rows = _table(wb["4 - Review queue"])

    assert headers == ["What", "Company", "Person", "LinkedIn", "CRM says", "Example says",
                       "Check"]
    kinds = [r["What"] for r in rows]
    # Most consequential first: a disputed verdict can hide a whole company.
    assert kinds[0] == "Profile fit disputed"
    assert "May have left" in kinds
    assert "Title differs" in kinds
    # The CRM holds both values and the provider has no record of the company:
    # nothing is missing, there is just no second opinion.
    assert "Company not found in Example" in kinds
    assert "Headcount or HQ missing" not in kinds
    assert kinds.index("Company not found in Example") > kinds.index("Profile fit disputed")
    initech = next(r for r in rows if r["What"] == "Company not found in Example")
    assert initech["Company"] == "Initech" and initech["LinkedIn"] in (None, "")
    assert "the only second opinion" in initech["Check"]
    moved = next(r for r in rows if r["What"] == "May have left" and r["Person"] == "Lee Park")
    assert moved["Example says"] == "now at Globex, as CMO (Example record updated 2026-06-01)"
    ray = next(r for r in rows if r["What"] == "May have left" and r["Person"] == "Ray Oh")
    assert "(matched on LinkedIn)" in ray["Example says"]
    # Each person row carries the CRM's LinkedIn URL, and the Check says so.
    assert ray["LinkedIn"] == "https://linkedin.com/in/ray-oh/"
    assert moved["LinkedIn"] in (None, "")                  # the CRM holds none for Lee
    assert moved["Check"] == "Open their LinkedIn (link in this row)"
    title_row = next(r for r in rows if r["What"] == "Title differs")
    assert title_row["LinkedIn"] == "https://www.linkedin.com/in/dana-reyes"
    assert title_row["Check"] == "Open their LinkedIn (link in this row)"
    # A handle match under another name becomes a question, not a finding.
    wrong = next(r for r in rows if r["What"] == "CRM LinkedIn may be someone else")
    assert wrong["Person"] == "Kim Lo" and "Kimberly Stone" in wrong["Example says"]
    assert wrong["LinkedIn"] == "https://www.linkedin.com/in/kimlo"
    # The move is the finding; a title at the new company is not a disagreement.
    assert not any(r["What"] == "Title differs" and r["Person"] == "Lee Park" for r in rows)
    # Every row says where to check.
    assert all(r["Check"] for r in rows)

    assert "Review queue" in html
    assert '<span class="fl">disputed — Example says yes</span>' in html
    assert '<span class="fl">no — now at Globex, as CMO</span>' in html


def test_one_crm_with_a_provider_still_compares_no_two_crms(tmp_path):
    """The one-CRM defect class, with a provider added: nothing comparing two
    CRMs appears when only one is configured."""
    wb, _, _ = _run(tmp_path, _Provider())
    for ws in wb.worksheets:
        headers, _ = _table(ws)
        assert "In HubSpot?" not in headers and "In Salesforce?" not in headers
        for row in ws.iter_rows(values_only=True):
            for v in row:
                assert "title only in" not in str(v or "")


def test_a_withheld_person_is_counted_shown_as_not_found_and_the_run_carries_on(tmp_path):
    """One record the provider will not show is one person not found — not a
    run with no artifact."""
    from quorum.enrich import Withheld

    class _Withholding(_Provider):
        def person_by_email(self, email):
            if email == "ari@acme.example":
                self.person_calls.append(email)
                raise Withheld("EXAMPLE_WITHHELD")
            return super().person_by_email(email)

    provider = _Withholding()
    cfg, sf, hs = _cfg(), _SF(), _HS()
    reconciled = [people_mod.reconcile(p, sf, hs) for p in _attendees()]
    pass_ = enrichment.start(provider)
    enrichment.not_in_crm(pass_, reconciled)
    enrichment.check_withheld(pass_)

    ari = next(r for r in reconciled if r.get("email") == "ari@acme.example")
    assert ari["other_found"] is False
    assert ari["other_name"] == "not found in Example"
    assert pass_.withheld == 1
    assert pass_.withheld_codes == {"EXAMPLE_WITHHELD"}
    # Everyone else not in the CRM was still looked up.
    assert len(provider.person_calls) == pass_.people_looked_up > 1


def test_every_lookup_withheld_is_the_provider_not_answering():
    from quorum.enrich import Withheld

    class _WithholdsAll(_Provider):
        def person_by_email(self, email):
            raise Withheld("")

    pass_ = enrichment.start(_WithholdsAll())
    cfg, sf, hs = _cfg(), _SF(), _HS()
    reconciled = [people_mod.reconcile(p, sf, hs) for p in _attendees()]
    enrichment.not_in_crm(pass_, reconciled)
    assert pass_.people_looked_up > 1
    with pytest.raises(enrichment.EveryLookupWithheld):
        enrichment.check_withheld(pass_)


def test_a_single_lookup_withheld_is_not_taken_for_an_outage():
    from quorum.enrich import Withheld

    class _WithholdsAll(_Provider):
        def person_by_email(self, email):
            raise Withheld("")

    pass_ = enrichment.start(_WithholdsAll())
    pass_.person("ari@acme.example")
    enrichment.check_withheld(pass_)  # one lookup cannot tell the two apart
    assert pass_.withheld == 1


def test_two_configured_providers_are_refused(monkeypatch):
    class _A:
        display_name, env_vars = "A", ("A_KEY",)

        @classmethod
        def from_env(cls, environ=None):
            return cls()

    class _B(_A):
        display_name, env_vars = "B", ("B_KEY",)

    monkeypatch.setattr(enrich, "_modules", lambda: iter([_A, _B]))
    with pytest.raises(enrich.MoreThanOneProvider):
        enrich.configured({})


# --- run.py wiring, against the real database -------------------------------- #


def test_the_weekly_run_checks_the_provider_first_and_adds_the_review_queue(
    database, gong_calls, tmp_path, monkeypatch
):
    """The order is the point: the provider's free check runs before the
    database is read or a CRM called, and the company pass before the map is
    chosen. No CRM is configured here (conftest), so every verdict is the
    CRM's "not assessed" beside the provider's.

    The config is a stub, not Config(): it borrows Config's two week methods
    and nothing else, so nothing is read from the environment.
    """
    from quorum.config import Config
    from quorum.weekly import run as run_mod
    from tests.test_import_and_weekly import (
        ACCOUNT, _import, _seed_account, _seed_profile,
    )

    account_id = _seed_account(database)
    _import(database, account_id, gong_calls)
    _seed_profile(database, account_id)

    provider = _Provider()
    provider.companies["acme.com"] = Company(name="Acme", domain="acme.com",
                                             employees=900, country="United States")
    events: list[str] = []
    provider.check = lambda: events.append("check") or "stub plan, 100 credits available"
    monkeypatch.setattr(run_mod.enrich, "configured", lambda: provider)
    real_connect = run_mod.db.connect
    monkeypatch.setattr(run_mod.db, "connect",
                        lambda cfg: (events.append("db"), real_connect(cfg))[1])

    cfg = SimpleNamespace(
        database_url=database, account=ACCOUNT, week_start="2026-08-17",
        output_dir=str(tmp_path), tz_offset="-04", shortlist_size=3, group_call_min=8,
        recent_days=90, customer_account_types=(), retain_runs=False,
        gong=SimpleNamespace(configured=False),
        salesforce=SimpleNamespace(
            configured=False, uses_client_credentials=False, access_token="",
            instance_url="", token_url="", client_id="", client_secret="",
            api_version="v61.0",
        ),
        hubspot=SimpleNamespace(configured=False, api_key=""),
    )
    cfg.week_bounds = lambda: Config.week_bounds(cfg)
    cfg.week_days_remaining = lambda: Config.week_days_remaining(cfg)
    logged: list[str] = []
    paths = run_mod.run_weekly(cfg, log=logged.append)

    assert events[:2] == ["check", "db"]
    assert any("Enrichment: Example" in line for line in logged)

    wb = load_workbook(paths["xlsx"])
    assert "4 - Review queue" in wb.sheetnames
    _, rows = _table(wb["2 - Company coverage"])
    acme = next(r for r in rows if r["Company"] == "acme.com")
    assert acme["Profile check"] == "CRM not assessed — Example says yes"
    # Not assessed is not disputed: there is no CRM verdict to disagree with.
    assert "(profile disputed)" not in json.dumps(
        [c.value for row in wb["3 - Stakeholder list"].iter_rows() for c in row], default=str
    )

    dump = json.loads(open(paths["json"]).read())
    assert dump["enrichment_provider"] == "Example"
    assert isinstance(dump["review_queue"], list)

    # The delivery step's copy names the provider too, and the window the
    # recent-contact count used.
    summary = json.loads(open(paths["summary"]).read())
    assert summary["enrichment_provider"] == "Example"
    assert summary["recent_days"] == 90


def test_a_provider_without_a_linkedin_lookup_is_not_asked(tmp_path):
    """The LinkedIn lookup is optional in the provider interface. Without it,
    an email miss stays a miss, and nothing else changes."""

    class _EmailOnly:
        display_name = "Example"

        def __init__(self):
            self._inner = _Provider()

        def person_by_email(self, email):
            return self._inner.person_by_email(email)

        def company_by_domain(self, domain):
            return self._inner.company_by_domain(domain)

    wb, _, _ = _run(tmp_path, _EmailOnly())
    _, rows = _table(wb["3 - Stakeholder list"])
    by = {r["Name"]: r for r in rows}
    assert by["Ray Oh"]["Still at company?"] == "not found in Example"
    _, queue = _table(wb["4 - Review queue"])
    assert not any(r["What"] == "CRM LinkedIn may be someone else" for r in queue)


@pytest.mark.parametrize(
    "a, b, agree",
    [
        ("Dana Reyes", "Dana Reyes", True),
        ("Dana M. Reyes", "dana reyes", True),       # middle initial, case
        ("José Álvarez", "Jose Alvarez", True),       # accents
        ("Mary-Kate O'Neil", "Mary Kate ONeil", True),   # hyphen, apostrophe
        ("Dana Reyes", "Dana Rivera", False),
        ("Dana Reyes", "Kimberly Stone", False),
        ("", "Dana Reyes", False),
    ],
)
def test_names_agree_on_first_and_last(a, b, agree):
    assert enrichment.names_agree(a, b) is agree


def test_the_stakeholder_list_says_how_it_is_built(tmp_path):
    """Which companies, which people, how many and in what order — in the
    reader's own values. A list of names with no stated rule reads as a
    recommendation from nowhere."""
    wb, html, _ = _run(tmp_path, _Provider())
    captions = [
        r[0] for r in wb["3 - Stakeholder list"].iter_rows(min_row=2, values_only=True)
        if r[0] and not any(r[1:])
    ]
    rule = captions[0]
    assert "50–5,000 employees, HQ in North America" in rule
    assert "plus any whose profile fit is disputed" in rule   # Globex is
    assert "up to 3 people from your CRM at VP or C-suite level" in rule
    assert "most senior first, then most recently contacted" in rule
    assert "People not in your CRM cannot appear here" in rule
    assert rule in html


# --- the move text, queue readability, and what the page claims -------------- #


def _mover(title):
    return Person(name="Sam Ito", title=title, employer_name="Hooli",
                  employer_domain="hooli.example",
                  current_jobs=(("hooli.example", "Hooli", title),))


class _OnePerson:
    display_name = "Example"

    def __init__(self, person):
        self.person = person

    def person_by_email(self, email):
        return self.person


@pytest.mark.parametrize(
    "title,expected",
    [("COO", "no — now at Hooli, as COO"), ("", "no — now at Hooli")],
)
def test_a_leavers_new_title_is_in_the_move_text_when_there_is_one(title, expected):
    row = {"name": "Sam Ito", "_email": "sam@acme.example", "domain": "acme.example",
           "company": "Acme", "title": "VP Ops"}
    enrichment.stakeholders(enrichment.start(_OnePerson(_mover(title))), [row])

    assert row["still_at"] == expected
    # The separate provider-title column stays blank for a mover.
    assert row["other_title"] == ""
    # The queue's "May have left" row reads it from `still_at`.
    queue = enrichment.review_queue(enrichment.start(_OnePerson(None)), [], [row])
    left = next(q for q in queue if q["kind"] == "May have left")
    assert left["other"] == expected[len("no — "):]


def test_a_company_the_crm_has_in_full_is_not_called_missing():
    """Missing is a gap in the CRM's data. A provider that does not know the
    company is a different finding, and the queue says so."""
    pass_ = enrichment.start(_OnePerson(None))
    nf = enrichment.not_found(pass_.provider)
    base = {"name": "Initech", "domain": "initech.example", "assessed": True,
            "disputed": False, "other_missing": True, "verdict_check": nf,
            "other_employees": None, "other_hq": ""}
    full = {**base, "employees": 300, "hq": "Canada", "meets": "yes"}
    gap = {**base, "domain": "umbrella.example", "name": "Umbrella", "employees": "",
           "hq": "Canada", "meets": "no size"}
    queue = enrichment.review_queue(pass_, [full, gap], [])

    assert [(q["kind"], q["company"]) for q in queue] == [
        ("Company not found in Example", "Initech"),
        ("Headcount or HQ missing", "Umbrella"),
    ]


def test_the_page_groups_the_queue_by_kind_and_the_sheet_does_not(tmp_path):
    wb, html, _ = _run(tmp_path, _Provider())
    _, rows = _table(wb["4 - Review queue"])
    section = html.split("<h2>Review queue", 1)[1].split("</section>", 1)[0]

    for kind in {r["What"] for r in rows}:
        n = sum(1 for r in rows if r["What"] == kind)
        assert f'<h3>{kind}<span class="cnt">{n}</span></h3>' in section
    # The kind is the sub-heading, so the column is dropped inside a group.
    assert "<th>What</th>" not in section and "<th>LinkedIn</th>" in section
    # The badge is the total, and the sheet stays one table.
    assert section.startswith(f'<span class="cnt">{len(rows)}</span>')
    assert [r["What"] for r in rows][0] == "Profile fit disputed"


def test_the_stakeholder_badge_counts_people_not_placeholders():
    from openpyxl import Workbook

    from quorum.weekly.stakeholders import NO_SENIOR_CONTACT

    ws = Workbook().active
    ws.append(["Company", "Name", "Title"])
    ws.append(["Acme", "Dana Reyes", "VP Sales"])
    ws.append(["Acme", "Lee Park", "VP Marketing"])
    ws.append(["Globex", NO_SENIOR_CONTACT, ""])
    ws.append(["Globex", "Kim Lo", "CRO"])

    html = view_mod.render_sheet(ws, "Stakeholder list")
    assert '<span class="cnt">3</span>' in html and html.count("<tr>") == 5


def test_the_page_wraps_cells_and_scrolls_tables_not_itself(tmp_path):
    _, html, _ = _run(tmp_path, _Provider())

    assert "text-overflow" not in html and "max-width:240px" not in html
    assert ".tw{overflow-x:auto}" in html
    assert html.count('<div class="tw"><table>') == html.count("<table>")


def _summary_sheet(rows):
    from openpyxl import Workbook

    ws = Workbook().active
    ws.append(["Tab", "What", "Count", "Out of"])
    for r in rows:
        ws.append(list(r))
    ws.append([])
    ws.append(["A footnote."])
    return ws


def test_the_summary_is_one_line_per_tab():
    ws = _summary_sheet([
        ("Company coverage", "Companies met", 8, None),
        ("Company coverage", "Meet your profile", 5, 8),
        ("Company coverage", "Could not be assessed", 0, 8),
        ("Met this week", "People", 20, None),
        ("Stakeholder list", "No title", 0, 12),
        ("Stakeholder list", "No mobile", 0, 12),
        ("Review queue", "Profile fit disputed", 0, None),
        ("Review queue", "May have left", 2, None),
        ("Review queue", "Title differs", 5, None),
        ("Review queue", "LinkedIn differs", 2, None),
    ])
    linkable = {"Company coverage", "Met this week", "Review queue"}
    out = view_mod.render_summary(ws, "Summary", linkable)
    text = dict(re.findall(r"<b>(?:<a[^>]*>)?([^<]+)(?:</a>)?</b> ([^<]*)</div>", out))

    # Every tab is present, and only the sheet's own words are used.
    assert list(text) == ["Company coverage", "Met this week", "Stakeholder list",
                          "Review queue"]
    # Zero rows are dropped; "of" appears only where Out of is set.
    assert text["Company coverage"] == "Companies met 8 · Meet your profile 5 of 8"
    assert "Could not be assessed" not in out
    assert text["Met this week"] == "People 20"
    # A tab that is all zero keeps one line, so it is not read as missing.
    assert text["Stakeholder list"] == "No title 0 of 12"
    assert "No mobile" not in out
    # The queue: the total, then the non-zero kinds, largest first, ties in
    # queue order.
    assert text["Review queue"] == (
        "9 to check: Title differs 5 · May have left 2 · LinkedIn differs 2"
    )
    assert "Profile fit disputed" not in out
    # Each tab name is a bold link to its section; the footnote stays.
    assert '<b><a href="#review-queue">Review queue</a></b>' in out
    assert '<b>Stakeholder list</b>' in out          # no such section: not a link
    assert out.rstrip().endswith('<p class="sub">A footnote.</p></section>')


def test_the_page_links_each_summary_tab_to_its_section(tmp_path):
    html = _run(tmp_path, _Provider())[1]  # no Summary sheet in this run
    assert 'id="review-queue"' in html and 'id="stakeholder-list"' in html


# --- two domains, and an email that is someone else's ------------------------ #


def _row(**kw):
    return {"name": "Sam Ito", "_email": "sam@acme.example", "domain": "acme.example",
            "company": "Acme", "title": "VP Ops", **kw}


def _still_at(person, **row):
    r = _row(**row)
    enrichment.stakeholders(enrichment.start(_OnePerson(person)), [r])
    return r["still_at"]


@pytest.mark.parametrize("jobs", [True, False])
def test_either_domain_counts_as_here(jobs):
    """Email from one domain, company on another: a position at the account's
    Website domain is still the same company."""
    at = ("acme.example", "Acme", "VP Ops")
    elsewhere = ("hooli.example", "Hooli", "COO")

    def person(job):
        return Person(name="Sam Ito", title=job[2], employer_name=job[1],
                      employer_domain=job[0], current_jobs=(job,) if jobs else ())

    two = {"domain": "acme.test", "account_domain": "acme.example"}
    assert _still_at(person(at), **two) == "yes"
    # Someone genuinely elsewhere still reads as having moved.
    assert _still_at(person(elsewhere), **two) == "no — now at Hooli, as COO"
    # The domain met still counts, and no account domain changes nothing.
    assert _still_at(person(at)) == "yes"
    assert _still_at(person(elsewhere), account_domain="") == "no — now at Hooli, as COO"


class _Lookups:
    display_name = "Example"

    def __init__(self, companies):
        self.companies, self.calls = companies, []

    def company_by_domain(self, domain):
        self.calls.append(domain)
        return self.companies.get(domain)


def _coverage(**kw):
    return [{"domain": "acme.test", "name": "Acme", "employees": 500, "hq": "US",
             "meets": "yes", "assessed": True, **kw}]


def test_the_company_lookup_falls_back_to_the_account_domain_once():
    found = Company(name="Acme", domain="acme.example", employees=520, country="United States")

    # Missed by the domain met, found by the account's: used, and said so.
    p = _Lookups({"acme.example": found})
    cov = _coverage(account_domain="acme.example")
    enrichment.companies(enrichment.start(p), cov, PROFILE)
    assert cov[0]["verdict_check"] == "agrees (looked up as acme.example)"
    assert cov[0]["other_employees"] == 520
    assert p.calls == ["acme.test", "acme.example"]

    # Found by neither: not found, and exactly one extra lookup.
    p = _Lookups({})
    cov = _coverage(account_domain="acme.example")
    enrichment.companies(enrichment.start(p), cov, PROFILE)
    assert cov[0]["verdict_check"] == "not found in Example"
    assert p.calls == ["acme.test", "acme.example"]

    # No account domain, or the same one: no extra lookup.
    for extra in ({}, {"account_domain": ""}, {"account_domain": "acme.test"}):
        p = _Lookups({})
        enrichment.companies(enrichment.start(p), _coverage(**extra), PROFILE)
        assert p.calls == ["acme.test"]

    # Found by the domain met: the account's is never asked.
    p = _Lookups({"acme.test": found})
    enrichment.companies(enrichment.start(p), _coverage(account_domain="acme.example"), PROFILE)
    assert p.calls == ["acme.test"]


class _ByEmailThenLinkedIn:
    display_name = "Example"

    def __init__(self, by_email, by_linkedin=None):
        self.by_email, self.by_linkedin, self.linkedin_calls = by_email, by_linkedin, []

    def person_by_email(self, email):
        return self.by_email

    def person_by_linkedin(self, url):
        self.linkedin_calls.append(url)
        return self.by_linkedin


def _colleague():
    return Person(name="Robin Hale", title="CRO", employer_name="Acme",
                  employer_domain="acme.example")


def test_an_email_match_under_another_name_is_not_used():
    provider = _ByEmailThenLinkedIn(_colleague())
    pass_ = enrichment.start(provider)
    r = _row(linkedin="https://www.linkedin.com/in/sam-ito")
    enrichment.stakeholders(pass_, [r])

    # Not "yes" and not the colleague's title: as if the email found nothing.
    assert r["still_at"] == "not found in Example"
    assert r["other_title"] == "" and r["matched_on"] == ""
    assert provider.linkedin_calls == ["https://www.linkedin.com/in/sam-ito"]
    assert pass_.email_name_rejected == 1

    queue = enrichment.review_queue(pass_, [], [r])
    assert [(q["kind"], q["who"], q["linkedin"], q["crm"], q["other"], q["check"])
            for q in queue] == [(
        "CRM email may belong to someone else", "Sam Ito",
        "https://www.linkedin.com/in/sam-ito", "sam@acme.example",
        "record at that email is Robin Hale",
        "Open their LinkedIn; the CRM email may be a colleague's",
    )]


def test_the_linkedin_fallback_can_still_find_the_right_person():
    sam = Person(name="Sam Ito", title="COO", employer_name="Hooli",
                 employer_domain="hooli.example")
    provider = _ByEmailThenLinkedIn(_colleague(), sam)
    r = _row(linkedin="https://www.linkedin.com/in/sam-ito")
    enrichment.stakeholders(enrichment.start(provider), [r])

    assert r["still_at"] == "no — now at Hooli, as COO (matched on LinkedIn)"
    assert r["email_other_person"] == "Robin Hale"


@pytest.mark.parametrize("name", ["Sam Q. Ito", "Sam Ito", ""])
def test_an_email_match_that_agrees_or_has_no_name_is_accepted_as_before(name):
    person = Person(name=name, title="VP Ops", employer_name="Acme",
                    employer_domain="acme.example")
    pass_ = enrichment.start(_OnePerson(person))
    r = _row()
    enrichment.stakeholders(pass_, [r])

    assert r["still_at"] == "yes" and r["matched_on"] == "email"
    assert "email_other_person" not in r and pass_.email_name_rejected == 0


def test_the_account_website_becomes_a_bare_domain():
    from quorum.crm.fieldmap import FieldMap
    from quorum.crm.salesforce import Salesforce
    from quorum.domains import bare_domain
    from tests.conftest import crm_config

    sf = Salesforce(crm_config(), FieldMap({}))
    sent = []
    sf.query = lambda soql: sent.append(soql) or {"records": [
        {"Name": "Acme", "Website": "https://www.acme.example/about"}]}

    assert sf.account_firmographics("001")["account_domain"] == "acme.example"
    assert "Website" in sent[0]
    sf.query = lambda soql: {"records": [{"Name": "Acme"}]}
    assert sf.account_firmographics("001")["account_domain"] == ""

    for raw, bare in [("HTTP://WWW.Acme.Example:8443/a?b=1#c", "acme.example"),
                      ("acme.example.", "acme.example"), ("www.acme.example", "acme.example"),
                      ("  ", ""), (None, "")]:
        assert bare_domain(raw) == bare


def test_with_hubspot_only_there_is_no_account_domain():
    from quorum.crm.fieldmap import FieldMap
    from quorum.crm.salesforce import Salesforce
    from tests.conftest import crm_config

    cfg = _cfg()
    cfg.salesforce.configured, cfg.hubspot.configured = False, True
    hs = _HS()
    hs.configured = True
    hs.count_domain = lambda d: 3
    cov = coverage_mod.build_coverage(
        cfg, people_mod.group_companies(_attendees()), PROFILE,
        Salesforce(crm_config(False), FieldMap({})), hs,
    )
    assert cov and all(c["account_domain"] == "" for c in cov)


# --- HTML tidy-ups ---------------------------------------------------------- #


def _queue_sheet(rows):
    from openpyxl import Workbook

    ws = Workbook().active
    ws.append(["What", "Company", "Person", "LinkedIn", "CRM says", "Example says", "Check"])
    for r in rows:
        ws.append(list(r))
    return ws


def test_a_group_drops_the_columns_that_are_empty_in_all_its_rows():
    ws = _queue_sheet([
        ("Profile fit disputed", "Globex", "", "", "10 employees", "150 employees", "Look"),
        ("Profile fit disputed", "Initech", "", "", "20 employees", "90 employees", "Look"),
        ("Title differs", "Acme", "Dana Reyes", "", "VP", "SVP", "Open"),
        ("Title differs", "Acme", "Lee Park", "https://www.linkedin.com/in/lee-park",
         "VP", "CMO", "Open"),
    ])
    out = view_mod.render_sheet(ws, "Review queue")
    disputed, title = out.split("<h3>")[1:]

    assert "<th>Person</th>" not in disputed and "<th>LinkedIn</th>" not in disputed
    assert "<th>Company</th>" in disputed and "<th>CRM says</th>" in disputed
    # One value is enough to keep a column, for every row of the group.
    assert "<th>Person</th>" in title and "<th>LinkedIn</th>" in title
    # Other sections keep their empty columns.
    plain = view_mod.render_sheet(ws, "Company coverage")
    assert "<th>Person</th>" in plain and "<th>LinkedIn</th>" in plain


def test_the_someone_else_row_does_not_repeat_the_url(tmp_path):
    wb, html, _ = _run(tmp_path, _Provider())
    _, rows = _table(wb["4 - Review queue"])
    wrong = next(r for r in rows if r["What"] == "CRM LinkedIn may be someone else")

    assert wrong["CRM says"] in (None, "")
    assert wrong["LinkedIn"] == "https://www.linkedin.com/in/kimlo"
    assert "Kimberly Stone" in wrong["Example says"] and wrong["Check"]
    group = html.split("<h3>CRM LinkedIn may be someone else")[1].split("<h3>")[0]
    assert "<th>CRM says</th>" not in group and "<th>LinkedIn</th>" in group


def _line(rows):
    out = view_mod.render_summary(_summary_sheet(rows), "Summary", {"Met this week"})
    return re.search(r"</b> ([^<]*)</div>", out).group(1)


def test_summary_labels_drop_a_lead_in_an_earlier_label_already_said():
    rows = [
        ("Met this week", "In your CRM", 6, 8),
        ("Met this week", "In your CRM, no title", 2, 6),
        ("Met this week", "In your CRM, no LinkedIn", 1, 6),
        ("Met this week", "In your CRM, no mobile", 3, 6),
    ]
    assert _line(rows) == "In your CRM 6 of 8 · no title 2 of 6 · no LinkedIn 1 of 6 · no mobile 3 of 6"

    # A hidden zero row still counts as an earlier label, and the sheet's own
    # order is what "earlier" means.
    for hidden in (1, 2):
        zeroed = [r if i != hidden else (*r[:2], 0, r[3]) for i, r in enumerate(rows)]
        text = _line(zeroed)
        assert text.startswith("In your CRM 6 of 8 · ")
        assert text.endswith("no mobile 3 of 6") and "In your CRM," not in text
    # Even with the plain label itself hidden: the others share its lead-in.
    assert "In your CRM," not in _line([(*rows[0][:2], 0, 8), *rows[1:]])


def test_summary_labels_keep_words_nothing_earlier_said():
    rows = [
        ("Met this week", "In your CRM", 6, 8),
        # Not followed by ", ": kept whole.
        ("Met this week", "In your CRM and found", 3, None),
        # Nothing earlier is exactly "People" or starts with "People, ".
        ("Met this week", "People, in all", 4, None),
        # The longest lead-in an earlier label supplies wins.
        ("Met this week", "In your CRM, no title", 2, 6),
        ("Met this week", "In your CRM, no title, and no mobile", 1, 6),
    ]
    assert _line(rows) == (
        "In your CRM 6 of 8 · In your CRM and found 3 · People, in all 4 · "
        "no title 2 of 6 · and no mobile 1 of 6"
    )
    # The sheet keeps its labels.
    ws = _summary_sheet(rows)
    view_mod.render_summary(ws, "Summary", set())
    assert ws["B6"].value == "In your CRM, no title, and no mobile"


# --- people met who are in the CRM with no title ---------------------------- #


def _titleless(**kw):
    return {"attendee_name": "Sam Ito", "crm_name": "Sam Ito", "email": "sam@acme.example",
            "domain": "acme.example", "in_salesforce": True, "in_hubspot": None,
            "title": "", "flag": "needs title", "linkedin_in_crm": True,
            "linkedin_url_in_crm": "https://www.linkedin.com/in/sam-ito", **kw}


class _Sam(_Provider):
    """The shared provider, plus someone the CRM holds without a title."""

    def __init__(self, person):
        super().__init__()
        self.people["sam@acme.example"] = person


def _run_with(tmp_path, provider, extra, enrich=True):
    cfg, sf, hs = _cfg(), _SF(), _HS()
    people = _attendees()
    reconciled = [people_mod.reconcile(p, sf, hs) for p in people] + extra
    coverage = coverage_mod.build_coverage(
        cfg, people_mod.group_companies(people), PROFILE, sf, hs
    )
    pass_ = enrichment.start(provider) if enrich else None
    queue = []
    if pass_:
        enrichment.companies(pass_, coverage, PROFILE)
    rows, raw = stakeholders_mod.build(
        cfg, coverage, coverage_mod.seniority_terms(PROFILE), {}, sf
    )
    if pass_:
        enrichment.stakeholders(pass_, rows)
        enrichment.not_in_crm(pass_, reconciled)
        enrichment.in_crm_no_title(pass_, reconciled, coverage)
        queue = enrichment.review_queue(pass_, coverage, rows, reconciled)
    from quorum.weekly import summary as summary_mod

    stats = summary_mod.build(
        cfg, reconciled, coverage, rows, raw, PROFILE,
        enrichment=provider.display_name if pass_ else None, queue=queue,
        queue_kinds=enrichment.queue_kinds(provider) if pass_ else (),
    )
    xlsx = str(tmp_path / "weekly_stakeholder_map_2026-08-17.xlsx")
    workbook_mod.build_workbook(
        cfg, reconciled, coverage, [], rows, xlsx, profile=PROFILE, geo_label="North America",
        enrichment=provider.display_name if pass_ else None, queue=queue, summary=stats,
    )
    return load_workbook(xlsx), pass_, rows, stats


def _sam_person(**kw):
    return Person(name="Sam Ito", title="VP Ops", employer_name="Acme",
                  employer_domain="acme.example", linkedin="https://www.linkedin.com/in/sam-ito-pro",
                  **kw)


def test_a_titleless_crm_record_is_looked_up_and_reported_not_listed(tmp_path):
    wb, pass_, rows, stats = _run_with(tmp_path, _Sam(_sam_person()), [_titleless()])
    _, met = _table(wb["1 - Met this week"])
    sam = next(r for r in met if r["Email"] == "sam@acme.example")

    assert (sam["Name (Example)"], sam["Title (Example)"], sam["LinkedIn (Example)"]) == (
        "Sam Ito", "VP Ops", "https://www.linkedin.com/in/sam-ito-pro")
    assert sam["LinkedIn (CRM)"] == "https://www.linkedin.com/in/sam-ito"
    _, queue = _table(wb["4 - Review queue"])
    row = next(r for r in queue if r["What"] == "Title missing in CRM")
    assert (row["Person"], row["CRM says"], row["Example says"], row["LinkedIn"]) == (
        "Sam Ito", "(none)", "VP Ops", "https://www.linkedin.com/in/sam-ito")
    assert row["Check"] == "Open their LinkedIn; fill Title in the CRM"
    assert sum(1 for r in queue if r["What"] == "Title missing in CRM") == 1
    # The list is still built from the CRM's Title: a provider's does not add anyone.
    assert not any(r["name"] == "Sam Ito" for r in rows)
    by = {s["key"]: (s["count"], s["out_of"], s["what"]) for s in stats}
    assert by["people_in_crm_no_title_found"] == (1, 1, "In your CRM, no title, and found by Example")
    assert by["queue:Title missing in CRM"][0] == 1
    assert pass_.titleless_looked_up == 1


def test_the_title_missing_row_falls_back_to_the_providers_url_and_names_a_move(tmp_path):
    away = Person(name="Sam Ito", title="COO", employer_name="Hooli",
                  employer_domain="hooli.example",
                  linkedin="https://www.linkedin.com/in/sam-ito-pro",
                  current_jobs=(("hooli.example", "Hooli", "COO"),))
    wb, *_ = _run_with(tmp_path, _Sam(away), [_titleless(linkedin_url_in_crm="")])
    _, queue = _table(wb["4 - Review queue"])
    row = next(r for r in queue if r["What"] == "Title missing in CRM")

    assert row["Example says"] == "now at Hooli, as COO"
    assert row["LinkedIn"] == "https://www.linkedin.com/in/sam-ito-pro"


def test_a_titleless_record_under_another_name_fills_nothing(tmp_path):
    other = Person(name="Robin Hale", title="CRO", employer_name="Acme",
                   employer_domain="acme.example")
    wb, pass_, _, stats = _run_with(tmp_path, _Sam(other), [_titleless()])
    _, met = _table(wb["1 - Met this week"])
    sam = next(r for r in met if r["Email"] == "sam@acme.example")
    _, queue = _table(wb["4 - Review queue"])

    assert not any(sam[h] for h in ("Name (Example)", "Title (Example)", "LinkedIn (Example)"))
    assert [r["What"] for r in queue].count("CRM email may belong to someone else") == 1
    assert "Title missing in CRM" not in [r["What"] for r in queue]
    assert pass_.email_name_rejected == 1
    assert {s["key"]: s["count"] for s in stats}["people_in_crm_no_title_found"] == 0


def test_a_shared_inbox_with_no_title_is_not_looked_up(tmp_path):
    provider = _Sam(_sam_person())
    inbox = _titleless(email="team@acme.example", flag="shared inbox — verify")
    _run_with(tmp_path, provider, [inbox])
    assert "team@acme.example" not in provider.person_calls


def test_with_no_provider_a_titleless_record_is_not_looked_up(tmp_path):
    provider = _Sam(_sam_person())
    wb, pass_, _, stats = _run_with(tmp_path, provider, [_titleless()], enrich=False)

    assert pass_ is None and provider.person_calls == []
    headers, _ = _table(wb["1 - Met this week"])
    assert not any("Example" in h for h in headers)
    assert "4 - Review queue" not in wb.sheetnames
    assert not any(s["key"] == "people_in_crm_no_title_found" for s in stats)


def test_the_linkedin_column_holds_the_url_a_dash_or_the_reason():
    from quorum.crm.fieldmap import NOT_AVAILABLE, NOT_CHECKED
    from quorum.weekly.workbook import NO_RECORD, _linkedin_cell

    url = "https://www.linkedin.com/in/sam-ito"
    assert _linkedin_cell(True, url) == url
    assert _linkedin_cell(False, "") == ""                    # a record with no URL
    assert _linkedin_cell(None) == NOT_AVAILABLE              # a CRM with no field
    assert _linkedin_cell(NOT_CHECKED) == NOT_CHECKED         # no Salesforce asked


@pytest.mark.parametrize("sf_on,hs_on", [(True, False), (False, True), (True, True), (False, False)])
def test_tab_1_linkedin_column_in_every_configured_state(tmp_path, sf_on, hs_on):
    from quorum.crm.fieldmap import NOT_CHECKED
    from quorum.weekly.workbook import NO_RECORD

    cfg = _cfg()
    cfg.salesforce.configured, cfg.hubspot.configured = sf_on, hs_on
    url = "https://www.linkedin.com/in/sam-ito"
    held = {"attendee_name": "Sam Ito", "email": "sam@acme.example", "domain": "acme.example",
            "in_salesforce": True if sf_on else None, "in_hubspot": True if hs_on else None,
            "title": "VP", "mobile_in_crm": True, "flag": "",
            "linkedin_in_crm": True if sf_on else NOT_CHECKED,
            "linkedin_url_in_crm": url if sf_on else ""}
    absent = {**held, "attendee_name": "Dana Reyes", "email": "dana@acme.example",
              "in_salesforce": False if sf_on else None, "in_hubspot": False if hs_on else None}
    xlsx = str(tmp_path / "weekly_stakeholder_map_2026-08-17.xlsx")
    workbook_mod.build_workbook(cfg, [held, absent], [], [], [], xlsx, profile=PROFILE,
                                geo_label="North America")
    headers, rows = _table(load_workbook(xlsx)["1 - Met this week"])
    by = {r["Name"]: r for r in rows}

    assert "LinkedIn?" not in headers
    if not (sf_on or hs_on):
        assert "LinkedIn (CRM)" not in headers
        return
    assert by["Sam Ito"]["LinkedIn (CRM)"] == (url if sf_on else NOT_CHECKED)
    assert by["Dana Reyes"]["LinkedIn (CRM)"] == NO_RECORD
