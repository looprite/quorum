"""The CRM's "no longer at the company" box.

Optional and resolved like the other logical fields. The rule these tests hold
is that *unresolved means exactly as before*: no column, no count, no exclusion.
A CRM without such a field gets nothing, not a false.

No real Config(); stubs, invented names, `.example` domains.
"""

from __future__ import annotations

import datetime as dt
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from quorum.crm.contact import Contact
from quorum.crm.fieldmap import FieldMap, resolve
from quorum.crm.salesforce import Salesforce
from quorum.weekly import enrichment
from quorum.weekly import people as people_mod
from quorum.weekly import stakeholders as stakeholders_mod
from quorum.weekly import summary as summary_mod
from quorum.weekly import workbook as workbook_mod
from tests.conftest import crm_config
from tests.test_fieldmap import ACCOUNT_FIELDS, CONTACT_FIELDS, FakeOrg, f

PROFILE = {"employee_count_min": 50, "employee_count_max": 5000,
           "hq_geographies": ["North America"], "focus_seniority": ["vp", "c-level"]}
TODAY = dt.date.today().isoformat()


def _cfg(sf=True, hs=False):
    return SimpleNamespace(
        salesforce=SimpleNamespace(configured=sf), hubspot=SimpleNamespace(configured=hs),
        customer_account_types=(), recent_days=90, group_call_min=8, shortlist_size=3,
    )


def person(name, left, title="VP Sales", activity=""):
    return Contact(name=name, title=title, email=f"{name.split()[0].lower()}@acme.example",
                   linkedin="", last_activity=activity, left_company=left)


COVERAGE = [{"domain": "acme.example", "name": "Acme", "assessed": True, "is_target": True,
             "met": 1, "employees": 500, "hq": "US", "meets": "yes"}]


class _SF:
    configured = True

    def __init__(self, bench):
        self.bench = bench

    def senior_bench(self, domain, terms):
        return list(self.bench)


def _build(bench, cfg=None):
    return stakeholders_mod.build(cfg or _cfg(), COVERAGE, ["vp"], {}, _SF(bench))


def _names(rows):
    return [r["name"] for r in rows]


# --- the field map ------------------------------------------------------------ #

LEFT_FIELDS = CONTACT_FIELDS + [
    f("No_Longer_At_Company__c", "No Longer at Company", "boolean"),
    f("Former_Employee__c", "Former Employee", "boolean"),
    f("Account_No_Longer_Customer__c", "Account No Longer Customer", "boolean"),
    f("Lead_Left_Company__c", "Lead Left Company", "boolean"),
    f("Former_Employee_Notes__c", "Former Employee Notes", "textarea"),
]


def test_the_field_resolves_like_the_others_and_decoys_are_rejected():
    org = FakeOrg(
        fields={"Account": ACCOUNT_FIELDS, "Contact": LEFT_FIELDS},
        populated={("Contact", "No_Longer_At_Company__c"): 900,
                   ("Contact", "Former_Employee__c"): 40,
                   ("Account", "NumberOfEmployees"): 10, ("Account", "BillingCountry"): 10},
    )
    field_map, prov = resolve(org, log=lambda *_: None)

    assert field_map["Contact"]["left_company"] == ["No_Longer_At_Company__c", "Former_Employee__c"]
    why = {r["field"]: r["why"] for r in
           prov["Contact"]["fields"]["left_company"]["rejected"]}
    assert why["Account_No_Longer_Customer__c"] == "excluded by pattern"
    assert why["Lead_Left_Company__c"] == "excluded by pattern"
    assert "cannot hold" in why["Former_Employee_Notes__c"]


def test_an_org_with_no_such_field_resolves_to_nothing_and_is_not_an_error():
    field_map, _ = resolve(FakeOrg(), log=lambda *_: None)
    assert field_map["Contact"]["left_company"] == []


def _adapter(names):
    return Salesforce(crm_config(), FieldMap({"Contact": {"left_company": names}}))


def test_the_adapter_reads_the_box_three_ways():
    record = {"Name": "Dana Reyes", "Email": "dana@acme.example", "No_Longer__c": True}
    assert _adapter(["No_Longer__c"])._contact(record).left_company is True
    assert _adapter(["No_Longer__c"])._contact({**record, "No_Longer__c": False}).left_company is False
    # Unresolved: None, which is not False and changes nothing.
    assert _adapter([])._contact(record).left_company is None
    sent = []
    sf = _adapter(["No_Longer__c"])
    sf.query = lambda soql: sent.append(soql) or {"records": []}
    sf.senior_bench("acme.example", ["VP"])
    assert "No_Longer__c" in sent[0]
    sf = _adapter([])
    sf.query = lambda soql: sent.append(soql) or {"records": []}
    sf.senior_bench("acme.example", ["VP"])
    assert "No_Longer" not in sent[1]


# --- the stakeholder list ------------------------------------------------------ #


def test_a_marked_contact_is_left_off_and_the_next_person_takes_the_place():
    bench = [person("Ada Lin", True), person("Ben Ode", False), person("Cy Pell", False),
             person("Di Quo", False), person("Ed Rau", False)]
    rows, raw = _build(bench)
    # Ada was the most senior of equals; the place goes to the next in the ranking.
    assert _names(rows) == ["Ben Ode", "Cy Pell", "Di Quo"]
    assert raw[0]["marked_left"] == 1

    stats = summary_mod.build(_cfg(), [], COVERAGE, rows, raw, PROFILE,
                              marked_left_field=True)
    by = {s["key"]: (s["count"], s["what"]) for s in stats}
    assert by["stakeholders_marked_left"] == (
        1, "Marked as left in your CRM, left off the list")


def test_no_leaver_row_for_someone_the_provider_says_moved():
    from quorum.enrich import Person

    class Provider:
        display_name = "Example"

        def person_by_email(self, email):
            name = "Ada Lin" if email.startswith("ada") else "Ben Ode"
            return Person(name=name, title="CTO", employer_name="Hooli",
                          employer_domain="hooli.example", updated="2026-06-01")

    rows, _ = _build([person("Ada Lin", True), person("Ben Ode", False)])
    pass_ = enrichment.start(Provider())
    enrichment.stakeholders(pass_, rows)
    queue = enrichment.review_queue(pass_, COVERAGE, rows)

    assert not any(q["who"] == "Ada Lin" for q in queue)
    assert [q["who"] for q in queue if q["kind"] == "May have left"] == ["Ben Ode"]


def test_a_marked_leavers_recent_activity_no_longer_counts():
    """The count of companies with a senior contact in the last 90 days."""
    leaver_only = [person("Ada Lin", True, activity=TODAY), person("Ben Ode", False)]
    rows, raw = _build(leaver_only)
    assert raw[0]["any_recent_contact"] is False
    stats = summary_mod.build(_cfg(), [], COVERAGE, rows, raw, PROFILE,
                              marked_left_field=True)
    recent = next(s for s in stats if s["key"] == "companies_fit_recent_senior")
    assert recent["count"] == 0

    # The same activity on a contact not marked as left does count.
    _, raw = _build([person("Ada Lin", False, activity=TODAY), person("Ben Ode", False)])
    assert raw[0]["any_recent_contact"] is True
    # And on a contact the CRM has no such field for.
    _, raw = _build([person("Ada Lin", None, activity=TODAY), person("Ben Ode", None)])
    assert raw[0]["any_recent_contact"] is True


def test_everyone_marked_leaves_a_stated_gap_not_an_empty_company():
    rows, raw = _build([person("Ada Lin", True), person("Ben Ode", True)])
    assert _names(rows) == [stakeholders_mod.NO_SENIOR_CONTACT]
    assert raw[0]["marked_left"] == 2


def test_resolved_false_changes_nothing():
    bench = [person("Ada Lin", False), person("Ben Ode", False), person("Cy Pell", False)]
    assert _build(bench)[0] == _build([person(n, None) for n in ("Ada Lin", "Ben Ode", "Cy Pell")])[0]
    assert "marked_left" not in _build(bench)[1][0]


def _book(tmp_path, bench, field):
    cfg = _cfg()
    rows, raw = _build(bench)
    stats = summary_mod.build(cfg, [], COVERAGE, rows, raw, PROFILE, marked_left_field=field)
    xlsx = str(tmp_path / f"weekly_stakeholder_map_2026-08-17-{field}.xlsx")
    workbook_mod.build_workbook(cfg, [], COVERAGE, [], rows, xlsx, profile=PROFILE,
                                geo_label="North America", summary=stats,
                                marked_left_field=field)
    wb = load_workbook(xlsx)
    return {ws.title: [r for r in ws.iter_rows(values_only=True)] for ws in wb}, stats


def test_unresolved_is_the_same_output_as_before(tmp_path):
    bench = [person("Ada Lin", None), person("Ben Ode", None)]
    off, stats_off = _book(tmp_path, bench, False)
    on, stats_on = _book(tmp_path, [person("Ada Lin", False), person("Ben Ode", False)], True)

    assert "stakeholders_marked_left" not in {s["key"] for s in stats_off}
    assert not any("no longer at the company" in str(c) for ws in off.values()
                   for r in ws for c in r)
    # Resolved adds exactly one Summary row and one caption; nothing else moves.
    assert [r for r in on["Summary"] if "Marked as left" not in str(r)] == off["Summary"]
    assert [r for r in on["3 - Stakeholder list"]
            if "no longer at the company" not in str(r)] == off["3 - Stakeholder list"]
    assert on["2 - Company coverage"] == off["2 - Company coverage"]
    assert any("no longer at the company" in str(r) for r in on["3 - Stakeholder list"])


# --- met this week ------------------------------------------------------------- #


class _Crm:
    def __init__(self, configured, contact=None):
        self.configured, self.contact = configured, contact
        self.linkedin_available = False

    def contact_by_email(self, email):
        return self.contact


@pytest.mark.parametrize("sf_on,hs_on", [(True, False), (False, True), (True, True), (False, False)])
def test_the_flag_in_every_configured_state(sf_on, hs_on):
    marked = Contact(name="Ada Lin", title="VP", email="ada@acme.example", left_company=True)
    # HubSpot has no such box: its record carries None.
    hubspot = Contact(name="Ada Lin", title="VP", email="ada@acme.example")
    sf = _Crm(sf_on, marked if sf_on else None)
    hs = _Crm(hs_on, hubspot if hs_on else None)
    r = people_mod.reconcile({"attendee_name": "Ada Lin", "email": "ada@acme.example",
                              "domain": "acme.example"}, sf, hs)

    assert ("marked as left in CRM" in r["flag"]) is sf_on
    # Nothing else changes for them: still in the CRM, title as read.
    if sf_on or hs_on:
        assert r["title"] == "VP" and "needs title" not in r["flag"]


def test_a_contact_not_marked_has_no_flag():
    for left in (False, None):
        sf = _Crm(True, Contact(name="Ada Lin", title="VP", email="a@acme.example",
                                left_company=left))
        r = people_mod.reconcile({"attendee_name": "Ada Lin", "email": "a@acme.example",
                                  "domain": "acme.example"}, sf, _Crm(False))
        assert "marked as left" not in r["flag"]
